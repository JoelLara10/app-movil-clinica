"""Build the Activity 16 hand-in from the student's retained cover."""
from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile
from datetime import datetime
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
EV = ROOT / 'evidencias'
REF = Path(r'C:\Users\JOEL LARA\Downloads\0. Portada  DSM.docx')
OUT = ROOT / 'Actividad 16 Autenticacion remota OftalCore.docx'
source_hash = sha256(REF.read_bytes()).hexdigest()
results = json.loads((EV / 'resultados-qa.json').read_text(encoding='utf-8'))

# Deterministic engineering diagram, not an illustration of an observed screen.
im = Image.new('RGB', (1600, 1420), 'white')
draw = ImageDraw.Draw(im)
font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 29)
small = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 25)
purple = '#66509a'
def text_center(x, y, lines, f=font):
    bb = draw.multiline_textbbox((0, 0), lines, font=f, spacing=6, align='center')
    draw.multiline_text((x-(bb[2]-bb[0])/2, y-(bb[3]-bb[1])/2), lines,
                        font=f, fill='#172039', spacing=6, align='center')
def box(x, y, w, h, text, fill='#f4f0fa', decision=False):
    if decision:
        draw.polygon([(x,y-h/2),(x+w/2,y),(x,y+h/2),(x-w/2,y)], fill=fill, outline=purple, width=3)
    else:
        draw.rounded_rectangle((x-w/2,y-h/2,x+w/2,y+h/2), radius=18, fill=fill, outline=purple, width=3)
    text_center(x,y,text)
def arrow(points, label=None, at=None):
    draw.line(points, fill=purple, width=4)
    x,y=points[-1]; px,py=points[-2]
    if y>py: tip=[(x,y),(x-9,y-15),(x+9,y-15)]
    elif y<py: tip=[(x,y),(x-9,y+15),(x+9,y+15)]
    elif x>px: tip=[(x,y),(x-15,y-9),(x-15,y+9)]
    else: tip=[(x,y),(x+15,y-9),(x+15,y+9)]
    draw.polygon(tip,fill=purple)
    if label: text_center(*at,label,small)
box(740,65,430,74,'Abrir la aplicación')
box(740,192,600,94,'Capturar usuario y contraseña')
box(740,322,600,98,'POST /auth/login mediante HTTPS')
box(740,452,600,98,'API valida cuenta activa\ny compara el hash bcrypt')
box(740,607,520,156,'¿Credenciales\nválidas?',decision=True)
box(1300,607,390,126,'HTTP 401\nMostrar rechazo\ny continuar en login',fill='#fff2f4')
box(740,782,600,106,'HTTP 200 y JWT\nGuardar sesión y actualizar contexto')
box(740,962,560,156,'¿Sesión y permisos\naceptados por la API?',decision=True)
box(260,962,390,126,'HTTP 401 o 403\nDenegar el recurso',fill='#fff2f4')
box(740,1132,600,100,'Acceder a la pantalla protegida\ny recursos permitidos')
box(740,1285,600,110,'Cerrar sesión\nRevocar JWT y borrar sesión local')
for start,end in [(102,145),(239,273),(371,403),(501,529),(685,729),(835,884),(1040,1082),(1182,1230)]:
    arrow([(740,start),(740,end)])
arrow([(1000,607),(1105,607)],'No',(1050,573))
arrow([(1300,544),(1300,192),(1040,192)])
text_center(786,710,'Sí',small)
arrow([(460,962),(455,962)],'No',(420,920))
text_center(786,1060,'Sí',small)
draw.text((80,1380),'La interfaz no sustituye la autorización que aplica el servidor.',font=small,fill='#444444')
im.save(EV / 'flujo-autenticacion.png')

d = Document(REF)
def replace_cover(index, text):
    p=d.paragraphs[index]
    p.runs[0].text=text
    for run in p.runs[1:]: run.text=''
replace_cover(10,'Actividad 16\nAutenticación remota y control de acceso en una aplicación móvil')
replace_cover(25,'28 de octubre de 2026')
body = d.add_section(WD_SECTION_START.NEW_PAGE)
body.page_width=Inches(8.5); body.page_height=Inches(11)
body.top_margin=Inches(.78); body.bottom_margin=Inches(.72)
body.left_margin=body.right_margin=Inches(.85)
body.header.is_linked_to_previous=False
body.footer.is_linked_to_previous=False
for name,size,bold in [('A16 Body',11,False),('A16 H1',15,True),('A16 H2',11.5,True),('A16 Caption',9.5,False),('A16 Small',9.5,False)]:
    st=d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st.font.name='Arial'; st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.line_spacing=1.10
    st.paragraph_format.space_after=Pt(7)
    if 'H' in name.split()[-1]: st.paragraph_format.keep_with_next=True
footer=body.footer.paragraphs[0]
footer.style=d.styles['A16 Small']; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('OftalCore   |   Actividad 16   |   ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text,style='A16 Body'):
    el=d.add_paragraph(text,style)
    return el
def h(text): return p(text,'A16 H1')
def sub(text): return p(text,'A16 H2')
def page(title): d.add_page_break();h(title)
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for i,s in enumerate(headers):t.rows[0].cells[i].text=s
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,s in enumerate(row): cells[i].text=str(s)
    for r,row in enumerate(t.rows):
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i])
            tcPr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'90');x.set(qn('w:type'),'dxa');mar.append(x)
            tcPr.append(mar)
            for par in c.paragraphs:
                par.style=d.styles['A16 Small'];par.paragraph_format.space_after=Pt(3)
                for run in par.runs:run.bold=(r==0)
            if r==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAEAEA');tcPr.append(shade)
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
    p('','A16 Small').paragraph_format.space_after=Pt(1)
    return t
def shots(a,b,caption):
    par=d.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(EV/a),width=Inches(2.62))
    par.add_run('     ')
    par.add_run().add_picture(str(EV/b),width=Inches(2.62))
    p(caption,'A16 Caption')

h('Introducción y análisis del caso')
p('OftalCore es una aplicación móvil de gestión hospitalaria desarrollada con React Native y Expo. Permite consultar módulos de pacientes, estudios, enfermería y monitoreo. Esta práctica utiliza su servicio remoto real de QA para demostrar cómo verificar la identidad del usuario y proteger el acceso a sus recursos. Se elige la modalidad c de la actividad, porque existe una API funcional.')
sub('Problema y justificación')
p('En el caso educativo, abrir una aplicación sin controles permitiría consultar calificaciones o datos de otros estudiantes. En OftalCore, el riesgo equivalente es consultar expedientes, estudios o signos vitales sin autorización. Los riesgos incluyen suplantación de identidad, divulgación de información y modificaciones no autorizadas. La solución valida las credenciales en la API por HTTPS y exige una sesión válida en las solicitudes protegidas.')
sub('Qué información debería estar protegida')
p('En el caso escolar deben protegerse calificaciones, matrícula, datos personales y trámites. En OftalCore se protegen los expedientes, estudios, datos de pacientes, alertas, cuentas de usuario y operaciones administrativas. También se protegen la contraseña y el token de sesión, porque permiten acceder al sistema.')
sub('Diferencia entre autenticación y autorización')
p('La autenticación comprueba quién es el usuario mediante sus credenciales. La autorización determina qué acciones y recursos puede utilizar. Un usuario puede iniciar sesión correctamente y recibir HTTP 403 si intenta acceder a un módulo que su rol no permite.')
sub('Qué envía la app al servicio remoto')
p('Envía el nombre de usuario y la contraseña en el cuerpo JSON de POST /auth/login mediante HTTPS. No envía un rol elegido por el cliente ni decide si la contraseña es correcta. Después utiliza el token recibido en el encabezado Authorization de las solicitudes protegidas.')
sub('Qué sucede con credenciales incorrectas')
p('La API devuelve HTTP 401 y un mensaje genérico; no emite una sesión válida. La app permanece en el formulario y muestra el rechazo. No debe revelar si falló el nombre de usuario o la contraseña. Un fallo de red se distingue de un rechazo de credenciales.')
p('La consigna exige realizar este análisis en la libreta. Las respuestas anteriores sirven para transcribirlo; la fotografía manuscrita debe incorporarse posteriormente como evidencia propia del estudiante.','A16 Small')

page('Diseño del flujo de autenticación')
p('El servidor controla el acceso; la navegación de la app refleja la sesión y los permisos. El siguiente diagrama resume el flujo implementado y la respuesta a los accesos rechazados.')
par=d.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
par.add_run().add_picture(str(EV/'flujo-autenticacion.png'),width=Inches(6.12))
p('Figura 1. Flujo de autenticación y control de acceso de OftalCore. Al volver a abrir la app con una sesión guardada, se confirma primero con GET /auth/me.','A16 Caption')
table(['Método y ruta','Función'],[
('POST /auth/login','Validar credenciales y emitir JWT'),
('GET /auth/me','Confirmar identidad y vigencia de sesión'),
('GET /catalog/services','Ejemplo de recurso protegido utilizado en la prueba'),
('POST /auth/logout','Revocar la sesión remota y limpiar la sesión local')],[2.45,4.25])

page('Implementación en el proyecto móvil')
p('Se reutilizó la aplicación existente y su API Flask conectada a MongoDB. No se creó una pantalla simulada ni otro backend. La configuración QA apunta a https://ineo-pbl-entrega.onrender.com/api/v1. El cliente de evidencia fue exportado desde este repositorio y se comunicó con ese servicio remoto.')
table(['Componente','Responsabilidad comprobada'],[
('LoginScreen.js','Captura campos, oculta contraseña, valida formato y muestra rechazo visible.'),
('AuthContext.js','Ejecuta login remoto; restaura sesión con /auth/me; gestiona salida.'),
('secureStorage.js','Guarda JWT en SecureStore nativo y elimina sesión al salir.'),
('api.js','Adjunta Bearer JWT y limpia almacenamiento ante HTTP 401.'),
('App.js y MainNavigator.js','Separan acceso público y módulos autenticados por rol.'),
('API auth y middleware','Verifican bcrypt, firma JWT, expiración, revocación y rol actual.')],[2.10,4.60])
sub('Validación de credenciales y sesión')
p('La API comprueba el estado activo de la cuenta y compara la contraseña con su hash bcrypt. Emite un JWT firmado con HS256 y expiración de ocho horas. Al verificarlo exige firma, expiración, emisor, audiencia y un identificador de sesión. Consulta el estado y el rol actuales en la base de datos, en lugar de confiar únicamente en datos enviados por la app.')
sub('Mejora incluida en esta actividad')
p('Se añadió un mensaje accesible dentro del formulario para mostrar errores de credenciales, campos o conexión también en Expo web. Se conservaron las alertas nativas y las traducciones existentes. La API mantiene la decisión de aceptar o rechazar el acceso.')
sub('Tres medidas de seguridad implementadas')
p('1. Transporte y validación remota. Las credenciales y los tokens viajan por HTTPS en QA y producción. La contraseña se valida mediante bcrypt en el servidor y no se conserva como contraseña de sesión en el móvil.')
p('2. Protección de la sesión. El JWT nativo se almacena con Expo SecureStore. Se verifica su vigencia al restaurar la sesión; el cierre revoca su identificador remoto y elimina datos locales. El JWT se utiliza en un encabezado, nunca como parámetro de URL.')
p('3. Autorización del servidor. Cada recurso protegido exige autenticación y, cuando corresponde, un rol permitido. La API responde 401 por sesión ausente o inválida y 403 por permisos insuficientes. Ocultar una opción del menú no reemplaza estos controles.')

page('Evidencia de acceso rechazado')
p('Capturas del cliente Expo web del proyecto móvil, en una ventana de 412 por 915 píxeles, conectado a QA. No son fotografías de una ejecución nativa Android. Se utiliza una cuenta de prueba autorizada; las contraseñas y los tokens no se publican.')
shots('01-inicio-sin-sesion.png','02-acceso-rechazado.png',
      'Figuras 2 y 3. Izquierda: al abrir sin sesión se presenta el formulario público. Derecha: una contraseña incorrecta produce el mensaje “Credenciales incorrectas” y mantiene al usuario en el acceso. El rechazo HTTP 401 está registrado en resultados-qa.json.')
p('Resultado observado: el intento rechazado no abrió el Dashboard ni generó acceso a los módulos protegidos. El campo de contraseña aparece enmascarado.','A16 Body')

page('Evidencia de acceso permitido y cierre')
p('La misma aplicación se probó con credenciales válidas contra la API remota. El servidor emitió el token y permitió acceder a la pantalla protegida. Después se cerró la sesión desde el Dashboard.')
shots('03-acceso-permitido.png','04-sesion-cerrada.png',
      'Figuras 4 y 5. Izquierda: Dashboard protegido disponible después de autenticar. Derecha: formulario de acceso recuperado al cerrar sesión. Los indicadores de la pantalla corresponden al entorno QA y no se utilizan como evidencia clínica.')
p('Resultado observado: login HTTP 200, identidad remota HTTP 200, recurso protegido HTTP 200 y logout HTTP 200. Al reutilizar el token revocado, /auth/me respondió HTTP 401.','A16 Body')

page('Pruebas y resultados obtenidos')
p('Las ocho pruebas siguientes se ejecutaron contra la API real de QA el 28 de octubre de 2026. El archivo resultados-qa.json conserva método, ruta, resultado y duración, sin contraseñas ni tokens. Los códigos esperados coincidieron con todos los obtenidos.')
rows=[]
for i,r in enumerate(results['pruebas'],1):
    rows.append((f'T{i:02}',r['prueba'],f"{r['metodo']} {r['ruta']}",str(r['obtenido'])))
table(['ID','Caso','Solicitud','HTTP'],rows,[.40,2.36,3.12,.82])
sub('Pruebas automatizadas complementarias')
p('Proyecto móvil: 52 pruebas aprobadas después del cambio del formulario. Se comprueban, entre otros aspectos, almacenamiento de token, limpieza de sesión, uso de encabezados y separación entre sesiones. Exportación Expo web QA completada correctamente.')
p('API aislada: 29 pruebas de autenticación y endurecimiento aprobadas. Una prueba adicional de permisos de monitoreo aprobó el rechazo HTTP 403 para roles no permitidos y HTTP 200 para un rol clínico autorizado. Esta prueba usa datos sintéticos y mongomock; no debe confundirse con una prueba de permisos ejecutada en QA remoto.')
sub('Alcance de la evidencia')
p('El documento acredita el flujo remoto y la interfaz Expo web del mismo código móvil. Las capturas no prueban por sí solas el almacenamiento nativo de Android. SecureStore se acredita mediante revisión de código y pruebas automatizadas. El APK nativo de QA existe, pero no se atribuyen aquí capturas de un teléfono que no se hayan registrado.')
p('El acceso por roles se evalúa también en la API. No se realizaron cambios de contraseña ni se alteraron expedientes para generar estas evidencias.','A16 Small')

page('Conclusión y referencias')
sub('Conclusión')
p('La autenticación remota permite verificar la identidad antes de entregar información. En OftalCore, la app envía credenciales por HTTPS y utiliza la sesión que autoriza la API. Las pruebas muestran acceso permitido con credenciales válidas y rechazo ante contraseña incorrecta, ausencia de token, token inválido y sesión revocada.')
p('La autorización completa esta protección al limitar las operaciones según el rol actual. Por eso la seguridad no depende únicamente de mostrar una pantalla de inicio de sesión: cada recurso debe validar la sesión y sus permisos en el servidor. El mismo principio se aplica al caso educativo para impedir que una persona consulte información académica ajena.')
sub('Fuentes técnicas consultadas')
for item in [
 'Expo. SecureStore. Almacenamiento protegido de valores en plataformas nativas. https://docs.expo.dev/versions/latest/sdk/securestore/',
 'React Navigation. Authentication flows. Separación de pantallas públicas y autenticadas. https://reactnavigation.org/docs/auth-flow/',
 'OWASP. Authentication Cheat Sheet. TLS, manejo de credenciales y mensajes de error. https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html',
 'PyJWT. Usage Examples. Verificación de firma y reclamaciones del token. https://pyjwt.readthedocs.io/en/stable/usage.html',
 'Código de la app móvil OftalCore. AuthContext, secureStorage, api, LoginScreen y navegación. https://github.com/JoelLara10/app-movil-clinica',
 'Código de la API. routes/auth.py, services/auth_service.py y middleware/auth_middleware.py. https://github.com/JoelLara10/ineo-pbl-entrega',
]: p(item,'A16 Small')
p('Consulta de fuentes y ejecución de evidencias: 28 de octubre de 2026. La documentación general de Expo se utiliza como referencia conceptual; el proyecto revisado utiliza Expo SDK 57.','A16 Small')
sub('Evidencia de la libreta')
p('Conforme a la instrucción de la actividad, el estudiante debe transcribir en la libreta las respuestas del análisis e insertar después una fotografía legible en este apartado. Este es el único elemento de entrega pendiente de incorporar.')

d.core_properties.title='Actividad 16 Autenticación remota y control de acceso en OftalCore móvil'
d.core_properties.subject='Desarrollo Móvil Integral'
d.core_properties.author='Lara Rodriguez Joel Alejandro'
d.core_properties.keywords='OftalCore, autenticación remota, autorización, aplicación móvil, JWT, HTTPS'
d.save(OUT)
from lxml import etree
from zipfile import ZIP_DEFLATED
with ZipFile(REF) as a, ZipFile(OUT) as b:
    originals = {n: a.read(n) for n in a.namelist() if n.startswith('word/media/') or n in (
        'word/header1.xml','word/footer1.xml','word/_rels/header1.xml.rels','word/_rels/footer1.xml.rels')}
    for n, content in originals.items():
        if content != b.read(n):
            assert etree.tostring(etree.fromstring(content), method='c14n') == etree.tostring(etree.fromstring(b.read(n)), method='c14n')
    temp = OUT.with_suffix('.tmp.docx')
    with ZipFile(temp, 'w', ZIP_DEFLATED) as fixed:
        for entry in b.infolist(): fixed.writestr(entry, originals.get(entry.filename, b.read(entry.filename)))
temp.replace(OUT)
with ZipFile(REF) as a,ZipFile(OUT) as b:
    preserved=[n for n in a.namelist() if n.startswith('word/media/') or n in ('word/header1.xml','word/footer1.xml','word/_rels/header1.xml.rels','word/_rels/footer1.xml.rels')]
    assert all(a.read(n)==b.read(n) for n in preserved),'Cambió un componente protegido de la portada'
assert sha256(REF.read_bytes()).hexdigest()==source_hash
print(str(OUT))
print('Portada original intacta; logos, encabezado, pie y medios conservados.')
